#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Утилиты для конвертера С2000-М"""

import os


def validate_file_path(filepath):
    """Валидация пути к файлу"""
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Файл не найден: {filepath}")
    if not os.path.isfile(filepath):
        raise ValueError(f"Не файл: {filepath}")
    return filepath


def check_missing_descriptions(ws, max_row):
    """Проверить отсутствующие описания.

    Красим только строки данных: столбец B заполнен (номер шлейфа/реле/
    выхода/считывателя), а столбец F пуст. Строки-разделители между
    приборами пропускаем.
    """
    missing = []
    for r in range(1, max_row + 1):
        a = ws.cell(row=r, column=1).value
        b = ws.cell(row=r, column=2).value
        f = ws.cell(row=r, column=6).value
        if isinstance(a, int):
            continue  # строка-заголовок прибора (Адрес)
        if b not in (None, '') and f in (None, ''):
            missing.append(r)
    return missing


def format_missing_cells(ws, missing_rows):
    """Пометить ячейки с отсутствующими описаниями"""
    from openpyxl.styles import Font, PatternFill
    
    bold_font = Font(name='Times New Roman', size=10, bold=True, italic=True)
    red_fill = PatternFill(fill_type='solid', fgColor='FF6347')
    
    for row_num in missing_rows:
        cell = ws.cell(row=row_num, column=6)
        cell.value = 'Где описание!!!???'
        cell.font = bold_font
        cell.fill = red_fill


def backup_file(filepath):
    """Создать бэкап файла перед перезаписью: <name>.bak.<YYYYmmdd_HHMMSS>"""
    import shutil
    import time
    if not os.path.exists(filepath):
        return None
    stamp = time.strftime('%Y%m%d_%H%M%S')
    bak = f"{filepath}.bak.{stamp}"
    shutil.copy2(filepath, bak)
    return bak


def validate_config_lines(lines):
    """Валидация строк конфигурации"""
    errors = []
    for i, line in enumerate(lines, 1):
        if line.strip() == '':
            continue
        if line.find(',') == -1:
            errors.append(f"Строка {i}: нет запятой в '{line.strip()[:50]}'")
    return errors


def normalize_line(line):
    """Нормализовать строку: убрать лишние пробелы, обработать описания"""
    line = line.strip()
    if not line:
        return line
    
    # Обработка описаний
    if 'Описание:' in line:
        idx = line.find('Описание:')
        line = line[idx:]
    
    # Обработка ключевых полей
    for prefix in ['Шлейф:', 'Раздел:', 'Время', 'Выход:', 'Реле:']:
        if prefix in line:
            idx = line.find(prefix)
            line = line[idx:]
            break
    
    return line


def safe_int(value, default=0):
    """Безопасно преобразовать в int"""
    try:
        return int(value)
    except (ValueError, TypeError):
        return default


def format_list_value(value):
    """Преобразовать список в строку без квадратных скобок и кавычек"""
    if isinstance(value, list):
        return str(value).strip('[]').strip('\'\'')
    return str(value)


if __name__ == '__main__':
    print("Утилиты готовы")
