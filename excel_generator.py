#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Генератор Excel для конфигурации С2000-М"""

import os
import platform
import subprocess
import sys as _sys
from typing import cast

try:
    from openpyxl import Workbook
except ImportError:
    subprocess.check_call([_sys.executable, '-m', 'pip', 'install', 'openpyxl', '-q', '--break-system-packages'])
    from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment, Color
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.worksheet import Worksheet


class ExcelGenerator:
    """Генератор Excel таблицы"""
    
    def __init__(self):
        self.wb = Workbook()
        self.ws = cast(Worksheet, self.wb.active)
        self.ws.title = 'Адреса, шлейфа'
        self.current_row = 1
        self.sections = {}
        self._setup_styles()
    
    def _setup_styles(self):
        """Настроить стили"""
#        self.thin = Side(border_style="hair")
        # openpyxl: color="000000" даёт rgb 00000000 (прозрачный) — нужен FF000000
        self.thin = Side(border_style="thin", color="FF000000")
        self.bold_font = Font(name='Times New Roman', size=10, bold=True)
        self.regular_font = Font(name='Times New Roman', size=10, italic=True)
        self.wrap = Alignment(wrap_text=True)
        self.light_fill = PatternFill(fill_type='solid', fgColor='F0F8FF')
        self.white_fill = PatternFill(fill_type='solid', fgColor='FFFAF0')
        self.header_fill = PatternFill(fill_type='solid', fgColor='FFFFE0')
        self.address_fill = PatternFill(fill_type='solid', fgColor='E6E6FA')
        self.error_fill = PatternFill(fill_type='solid', fgColor='FF6347')
    
    def set_border(self, cell_range='A1:F1'):
        """Установить границы ячеек"""
        for row in self.ws[cell_range]:
            for cell in row:
                cell.border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
    
    def add_title(self, list_str, row=1, column=1):
        """Добавить заголовок"""
        self.ws.merge_cells('A1:C1')
        for cel in list_str:
            if cel.find('Конфигурация ') != -1:
                column = 1
            if cel.find('Версия:') != -1:
                column = 4
            cell = self.ws.cell(row=row, column=column, value=cel)
            cell.alignment = self.wrap
            cell.font = self.bold_font
            cell.fill = self.light_fill if row % 2 == 0 else self.white_fill
            column += 1
        row += 1
        return row
    
    def add_address_row(self, list_str, row=1, column=1):
        """Добавить строку адреса"""
        title_sh = ['', 'Шлейф', 'Раздел', 'Тип шлейфа', '', 'Описание']
        
        for cel in list_str:
            if cel.find('Адрес:') != -1:
                column = 1
            if cel.find('Тип_прибора: ') != -1:
                cell_value = cel[cel.find('Тип_прибора: '):].replace('Тип_прибора: ', '')
                try:
                    device_id = int(cell_value)
                    cel = str(self._config.get_device_type(device_id))
                except (ValueError, AttributeError):
                    pass
            if column == 2 and cel.find('Сценарий_упр:') == -1:
                t_cell = ' '
                cell = self.ws.cell(row=row, column=3, value=t_cell)
                cell.border = Border(bottom=self.thin, top=self.thin, left=self.thin, right=self.thin)
                cell.fill = self.address_fill
            if cel.find('Сценарий_упр:') != -1:
                column = 3
            if cel.find('Версия:') != -1:
                column = 4
            if cel.find('Описание:') != -1:
                cel = cel.replace('Описание:', '').replace('\"', '')
                column = 6
            cell = self.ws.cell(row=row, column=column, value=cel)
            cell.font = self.bold_font
            cell.border = Border(bottom=self.thin, top=self.thin, left=self.thin, right=self.thin)
            cell.fill = self.address_fill
            column += 1
        
        row += 1
        column = 1
        # Заполняем заголовок
        for it in title_sh:
            cell = self.ws.cell(row=row, column=column, value=it)
            cell.font = self.bold_font
            cell.border = Border(bottom=self.thin, top=self.thin)
            cell.fill = self.header_fill
            column += 1
        row += 1
        return row
    
    def add_output_row(self, list_str, row=1, column=1):
        """Добавить строку шлейфа/выхода/реле"""
        group_start = row
        level = 1
        
        for cel in list_str:
            if cel.find('Шлейф:') != -1:
                column = 2
                level = 1
            elif cel.find('Выход:') != -1:
                column = 2
                level = 2
            elif cel.find('Реле:') != -1:
                column = 2
                level = 3
                self.ws.row_dimensions[row].height = 25
            elif cel.find('Раздел:') != -1:
                try:
                    int_r = int(cel.replace(' ', '').replace('Раздел:', ''))
                    if self._config:
                        section = self._config.get_section(int_r)
                        if section:
                            cel = section['description']
                        else:
                            cel = ''
                except (ValueError, AttributeError):
                    pass
                self.ws.row_dimensions[row].height = 25
                column = 3
            elif cel.find('Программа:') != -1 and str(self.ws.cell(row=row, column=2).value).find('Реле:') != -1:
                try:
                    int_relay = int(cel.replace(' ', '').replace('Программа:', ''))
                    cel = 'Пр. упр: ' + str(self._config.get_relay(int_relay))
                except (ValueError, AttributeError):
                    pass
                if 20 < len(cel) < 30:
                    self.ws.row_dimensions[row].height = 25
                if len(cel) > 30:
                    self.ws.row_dimensions[row].height = 37
                column = 3
            elif cel.find('Тип_шлейфа:') != -1:
                try:
                    int_cable = int(cel.replace(' ', '').replace('Тип_шлейфа:', ''))
                    cel = str(self._config.get_cable_type(int_cable))
                except (ValueError, AttributeError):
                    pass
                self.ws.row_dimensions[row].height = 25
                column = 4
            elif cel.find('Описание:') != -1:
                cel = cel.replace('Описание:', '')
                column = 6
            elif cel.find('Время') != -1:
                column = 4
            
            cell = self.ws.cell(row=row, column=column, value=cel)
            cell.font = self.regular_font
            cell.alignment = self.wrap
            cell.fill = self.light_fill if row % 2 == 0 else self.white_fill
            column += 1
        
        self.ws.row_dimensions.group(start=group_start, end=row, outline_level=level, hidden=False)
        row += 1
        return row
    
    def set_config(self, config):
        """Установить конфигурацию для доступа к методам"""
        self._config = config
    
    def set_sections(self, sections):
        """Установить разделы"""
        self.sections = sections
    
    def save(self, filepath):
        """Сохранить файл"""
        self.wb.save(filepath)
    
    def open_file(self, filepath):
        """Открыть файл (Linux/Windows/macOS)"""
        system = platform.system()
        try:
            if system == 'Linux':
                subprocess.call(['xdg-open', filepath])
            elif system == 'Darwin':  # macOS
                subprocess.call(['open', filepath])
            else:  # Windows
                os.startfile(filepath)
        except Exception as e:
            print(f"Не удалось открыть файл: {e}", file=sys.stderr)
    
    def mark_missing_descriptions(self, max_row):
        """Пометить ячейки с отсутствующими описаниями"""
        for row in self.ws.iter_rows(min_row=1, min_col=6, max_col=6, max_row=max_row):
            for cell in row:
                if cell.value is None:
                    cell.fill = self.error_fill
                    cell.value = 'Где описание!!!???'
                    cell.font = self.bold_font
    
    def set_columns_width(self):
        """Установить ширину столбцов"""
        self.ws.column_dimensions['A'].width = 10
        self.ws.column_dimensions['B'].width = 13
        self.ws.column_dimensions['C'].width = 18
        self.ws.column_dimensions['D'].width = 20
        self.ws.column_dimensions['E'].width = 18
        self.ws.column_dimensions['F'].width = 25
    
    def hide_columns(self, hidden_cols):
        """Скрыть столбцы"""
        for col in hidden_cols:
            self.ws.column_dimensions[col].hidden = True
    
    def set_devices(self, devices):
        """Установить список приборов для генерации"""
        self._devices = devices
    
    def generate(self, devices=None):
        """Сгенерировать Excel из структурированных данных приборов"""
        if devices is None:
            devices = getattr(self, '_devices', [])
        
        row = 1
        for dev in devices:
            row = self._add_device_block(dev, row)
        
        # ИСПРАВЛЕНО: Вместо "thick" используем объект границы на основе вашего self.thin
        # Либо задаем тонкую черную линию напрямую: Side(border_style="thin", color="000000")
        fine_border = Border(
            top=self.thin,
            bottom=self.thin,
            left=self.thin,
            right=self.thin
        )
        
        # Применяем тонкую границу ко всем заполненным ячейкам таблицы
        # (включая 6-й столбец с описанием); строки-разделители пропускаем
        for r in range(1, self.ws.max_row + 1):
            a = self.ws.cell(row=r, column=1).value
            b = self.ws.cell(row=r, column=2).value
            if a in (None, '') and b in (None, ''):
                continue  # пустая строка-разделитель
            for c in range(1, 7):
                self.ws.cell(row=r, column=c).border = fine_border
        
        return row
    
    def _add_device_block(self, dev, row):
        """Добавить блок одного прибора"""
        # Строка адреса
        row = self._add_address_header(dev, row)
        
        # Шлейфы
        for shleif in dev.shleifs:
            row = self._add_shleif_row(shleif, row)
        
        # Реле/выходы объединены по номеру
        # Сначала соберем все реле и добавим к ним описания выходов;
        # выходы, не совпавшие ни с одним реле, остаются standalone
        merged_relays = []
        output_by_num = {}
        for o in dev.outputs:
            output_by_num[o['num']] = o
        consumed_outputs = set()
        for r in dev.relays:
            if not (r.get('program_id') and r['program_id'] != 0):
                continue  # реле не отобразится — выход остаётся standalone
            merged = dict(r)
            if r['num'] in output_by_num:
                out = output_by_num[r['num']]
                if out.get('description'):
                    merged['output_desc'] = out['description']
                if out.get('section_id') and not merged.get('section_id'):
                    merged['section_id'] = out['section_id']
                consumed_outputs.add(r['num'])
            merged_relays.append(merged)

        for relay in merged_relays:
            row = self._add_relay_row(relay, row)

        # Стендаун-выходы (нет реле с тем же номером) — раньше терялись
        for num in sorted(n for n in output_by_num if n not in consumed_outputs):
            row = self._add_output_row(output_by_num[num], row)
        
        # Считыватели
        for reader in dev.readers:
            row = self._add_reader_row(reader, row)
        
        row += 1  # пустая строка между приборами
        return row
    
    def _add_address_header(self, dev, row):
        """Добавить строку адреса и заголовок"""
        title_sh = ['', 'Шлейф', 'Раздел', 'Тип шлейфа', 'Программа', 'Описание']
        
        # Строка с адресом
        self.ws.cell(row=row, column=1, value=dev.addr)
        self.ws.cell(row=row, column=1).font = self.bold_font
        self.ws.cell(row=row, column=1).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=1).fill = self.address_fill
        
        self.ws.cell(row=row, column=2, value=dev.type_name)
        self.ws.cell(row=row, column=2).font = self.bold_font
        self.ws.cell(row=row, column=2).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=2).fill = self.address_fill
        
        self.ws.cell(row=row, column=3, value='')
        self.ws.cell(row=row, column=3).font = self.bold_font
        self.ws.cell(row=row, column=3).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=3).fill = self.address_fill
        
        self.ws.cell(row=row, column=4, value=dev.version)
        self.ws.cell(row=row, column=4).font = self.bold_font
        self.ws.cell(row=row, column=4).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=4).fill = self.address_fill
        
        # Столбец E (пустой)
        self.ws.cell(row=row, column=5, value='')
        self.ws.cell(row=row, column=5).font = self.bold_font
        self.ws.cell(row=row, column=5).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=5).fill = self.address_fill
        
        self.ws.cell(row=row, column=6, value=dev.description)
        self.ws.cell(row=row, column=6).font = self.bold_font
        self.ws.cell(row=row, column=6).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=6).fill = self.address_fill
        self.ws.cell(row=row, column=6).alignment = self.wrap
        
        row += 1
        
        # Заголовок
        for i, title in enumerate(title_sh, 1):
            cell = self.ws.cell(row=row, column=i, value=title)
            cell.font = self.bold_font
            cell.border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
            cell.fill = self.header_fill
        self.ws.column_dimensions['E'].width = 20  # Ширина для программы
        row += 1
        return row
    
    def _add_shleif_row(self, shleif, row):
        """Добавить строку шлейфа. Только подключённые (с section_id) или с описанием."""
        has_section = bool(shleif.get('section_id'))
        has_desc = bool(shleif.get('description'))
        if not has_section and not has_desc:
            return row  # Нет ни раздела, ни описания — мусор
        self.ws.cell(row=row, column=2, value=shleif['num'])
        self.ws.cell(row=row, column=2).font = self.regular_font
        self.ws.cell(row=row, column=2).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=2).alignment = self.wrap
        self.ws.cell(row=row, column=2).fill = self.light_fill if row % 2 == 0 else self.white_fill
        
        # Столбец A (номер прибора) — дублируем border
        self.ws.cell(row=row, column=1, value='')
        self.ws.cell(row=row, column=1).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        
        section_desc = ''
        if shleif.get('section_id'):
            section = self.sections.get(shleif['section_id'])
            if section:
                section_desc = section['description']
        
        self.ws.cell(row=row, column=3, value=section_desc)
        self.ws.cell(row=row, column=3).font = self.regular_font
        self.ws.cell(row=row, column=3).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=3).alignment = self.wrap
        
        self.ws.cell(row=row, column=4, value=shleif.get('type_name', ''))
        self.ws.cell(row=row, column=4).font = self.regular_font
        self.ws.cell(row=row, column=4).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=4).alignment = self.wrap
        
        # Столбец E (Программа) — пустой для шлейфов
        self.ws.cell(row=row, column=5, value='')
        self.ws.cell(row=row, column=5).font = self.regular_font
        self.ws.cell(row=row, column=5).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=5).alignment = self.wrap
        
        self.ws.cell(row=row, column=6, value=shleif.get('description', ''))
        self.ws.cell(row=row, column=6).font = self.regular_font
        self.ws.cell(row=row, column=6).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=6).alignment = self.wrap
        
        row += 1
        return row
    
    def _add_relay_row(self, relay, row):
        """Добавить строку реле. Только реле с программой (program_id != 0)."""
        if not relay.get('program_id') or relay['program_id'] == 0:
            return row  # Пропускаем реле без программы
        self.ws.cell(row=row, column=2, value=f'Рел.{relay["num"]}')
        self.ws.cell(row=row, column=2).font = self.regular_font
        self.ws.cell(row=row, column=2).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=2).alignment = self.wrap
        
        self.ws.row_dimensions[row].height = 25
        
        # Раздел — из section_id реле
        section_desc = ''
        if relay.get('section_id'):
            section = self.sections.get(relay['section_id'])
            if section:
                section_desc = section['description']
        
        self.ws.cell(row=row, column=3, value=section_desc)
        self.ws.cell(row=row, column=3).font = self.regular_font
        self.ws.cell(row=row, column=3).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=3).alignment = self.wrap
        
        # Тип шлейфа — пусто
        self.ws.cell(row=row, column=4, value='')
        self.ws.cell(row=row, column=4).font = self.regular_font
        self.ws.cell(row=row, column=4).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=4).alignment = self.wrap
        
        # Программа (столбец 5)
        program_name = relay.get('program_name', '')
        if program_name:
            desc = f"Пр. упр: {program_name}"
        else:
            desc = relay.get('description', '')
        
        self.ws.cell(row=row, column=5, value=desc)
        self.ws.cell(row=row, column=5).font = self.regular_font
        self.ws.cell(row=row, column=5).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=5).alignment = self.wrap
        if len(desc) > 30:
            self.ws.row_dimensions[row].height = 37
        
        # Описание (столбец 6) — выход описания优先, иначе реле
        output_desc = relay.get('output_desc', '')
        relay_desc = relay.get('description', '')
        desc_value = output_desc if output_desc else relay_desc
        
        self.ws.cell(row=row, column=6, value=desc_value)
        self.ws.cell(row=row, column=6).font = self.regular_font
        self.ws.cell(row=row, column=6).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=6).alignment = self.wrap
        
        row += 1
        return row

    def _add_output_row(self, output, row):
        """Добавить строку выхода. Только с описанием или section_id."""
        if not output.get('description') and not output.get('section_id'):
            return row  # Пропускаем выходы без описания и раздела
        self.ws.cell(row=row, column=2, value=f'Вых.{output["num"]}')
        self.ws.cell(row=row, column=2).font = self.regular_font
        self.ws.cell(row=row, column=2).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=2).alignment = self.wrap
        
        section_desc = ''
        if output.get('section_id'):
            section = self.sections.get(output['section_id'])
            if section:
                section_desc = section['description']
        
        self.ws.cell(row=row, column=3, value=section_desc)
        self.ws.cell(row=row, column=3).font = self.regular_font
        self.ws.cell(row=row, column=3).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=3).alignment = self.wrap
        
        self.ws.cell(row=row, column=6, value=output.get('description', ''))
        self.ws.cell(row=row, column=6).font = self.regular_font
        self.ws.cell(row=row, column=6).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=6).alignment = self.wrap
        
        row += 1
        return row
    
    def _add_reader_row(self, reader, row):
        """Добавить строку считывателя. Только с section_id ИЛИ description."""
        has_section = bool(reader.get('section_id'))
        has_desc = bool(reader.get('description'))
        if not has_section and not has_desc:
            return row  # Нет ни раздела, ни описания — мусор
        self.ws.cell(row=row, column=2, value=reader['num'])
        self.ws.cell(row=row, column=2).font = self.regular_font
        self.ws.cell(row=row, column=2).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=2).alignment = self.wrap
        
        section_desc = ''
        if reader.get('section_id'):
            section = self.sections.get(reader['section_id'])
            if section:
                section_desc = section['description']
        
        self.ws.cell(row=row, column=3, value=section_desc)
        self.ws.cell(row=row, column=3).font = self.regular_font
        self.ws.cell(row=row, column=3).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=3).alignment = self.wrap
        
        self.ws.cell(row=row, column=6, value=reader.get('description', ''))
        self.ws.cell(row=row, column=6).font = self.regular_font
        self.ws.cell(row=row, column=6).border = Border(top=self.thin, bottom=self.thin, left=self.thin, right=self.thin)
        self.ws.cell(row=row, column=6).alignment = self.wrap
        
        row += 1
        return row
    
    def set_printer_settings(self, paper_size, orientation):
        """Установить настройки принтера"""
        self.ws.set_printer_settings(paper_size, orientation)
    
    def set_page_breaks(self):
        """Установить разрывы страниц"""
        self.ws.col_breaks.brk = [Break(6)]
    
    def set_view_mode(self, mode):
        """Установить режим просмотра"""
        self.ws.sheet_view.view = mode


if __name__ == '__main__':
    # Тест
    gen = ExcelGenerator()
    gen.set_printer_settings(gen.wb.PAPERSIZE_A4, gen.wb.ORIENTATION_PORTRAIT)
    gen.set_page_breaks()
    gen.set_view_mode('pageBreakPreview')
    gen.hide_columns(['E'])
    gen.set_columns_width()
    gen.save('test_output.xlsx')
    print("Тестовый файл создан")
